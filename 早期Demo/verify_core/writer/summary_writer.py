"""标识方式②：把全部差异项按 1、2、3 顺序汇总成独立 docx，供审批人确认。"""
from __future__ import annotations

from pathlib import Path

from ..models import VerifyReport, VerifyType

_TYPE_LABEL = {
    VerifyType.MISSING_IN_REPORT: "报告草稿缺失",
    VerifyType.EXTRA_IN_REPORT: "报告草稿多出",
    VerifyType.FIELD_MISMATCH: "内容不一致",
}


def write_summary(report: VerifyReport, out_path: str, title: str = "报告草稿核对差异汇总") -> str:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, level=0)

    counts = report.counts()
    doc.add_paragraph(
        f"海信资料行数：{report.a_rows}　报告草稿行数：{report.b_rows}　"
        f"差异项：{len(report.diffs)}（缺失 {counts.get('missing', 0)} / "
        f"多出 {counts.get('extra', 0)} / 不一致 {counts.get('mismatch', 0)}）"
    )

    diffs = report.diffs
    if not diffs:
        doc.add_paragraph("未发现差异，报告草稿与海信资料一致。")
        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(out_path)
        return out_path

    # 差异明细表：序号 / 类型 / 块 / 名称 / 型号 / 海信资料 / 报告草稿 / 说明
    doc.add_heading("差异项明细", level=1)
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(["序号", "类型", "清单", "关键件名称", "型号", "海信资料", "报告草稿", "说明"]):
        hdr[i].text = h

    for idx, it in enumerate(diffs, start=1):
        a, b = it.a_row, it.b_row
        name = (a.关键件名称 if a else "") or (b.关键件名称 if b else "")
        model = (a.型号 if a else "") or (b.型号 if b else "")
        # 差异字段的两侧取值（仅 FIELD_MISMATCH 有意义）
        a_val = b_val = ""
        if it.type == VerifyType.FIELD_MISMATCH:
            a_val = "；".join(f"{f}={getattr(a, f, '')}" for f in it.changed_fields)
            b_val = "；".join(f"{f}={getattr(b, f, '')}" for f in it.changed_fields)
        cells = table.add_row().cells
        vals = [
            str(idx),
            _TYPE_LABEL.get(it.type, it.type.value),
            "安全" if it.block == "safety" else "EMC",
            name, model, a_val, b_val, it.describe(),
        ]
        for i, v in enumerate(vals):
            cells[i].text = v

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
