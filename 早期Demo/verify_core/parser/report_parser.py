"""文档解析入口。

- parse_source_list：解析海信认证资料(A)——纯关键件清单文档
- parse_report_draft：解析认证机构报告草稿(B)——从 29+ 张混合表中
  用表头指纹**定位并抽取**关键件清单表，忽略照片/描述等无关表。

两者都返回 (rows, docx_path)。docx_path 供 comment_writer 二次打开写批注。
"""
from __future__ import annotations

from pathlib import Path

from ..models import Row
from .converter import ensure_docx
from .table_parser import is_keylist_header, parse_keylist_table


def _guess_block(rows_seen: int, table_index: int) -> str:
    """粗略推断表属于安全还是 EMC 清单。

    初稿策略：报告草稿中安全清单通常先出现、体量大；EMC 清单靠后。
    这里按"首个命中的大表=safety，其后命中的=emc"简单分派，后续可用
    上下文段落标题（"电磁兼容关键件清单"）精确化。
    """
    return "safety" if rows_seen == 0 else "emc"


def _open(path: str):
    from docx import Document

    docx_path = ensure_docx(path)
    return Document(str(Path(docx_path).resolve())), docx_path


def parse_source_list(path: str, side: str = "A") -> tuple[list[Row], str]:
    """解析海信认证资料：文档内所有关键件清单表都要。"""
    document, docx_path = _open(path)
    rows: list[Row] = []
    block_seq = 0
    for idx, table in enumerate(document.tables, start=1):
        if not is_keylist_header(table):
            continue
        block = _guess_block(block_seq, idx)
        rows.extend(parse_keylist_table(table, idx, side=side, block=block))
        block_seq += 1
    return rows, docx_path


def parse_report_draft(path: str, side: str = "B") -> tuple[list[Row], str]:
    """解析认证机构报告草稿：只抽取命中表头指纹的关键件清单表。"""
    document, docx_path = _open(path)
    rows: list[Row] = []
    block_seq = 0
    for idx, table in enumerate(document.tables, start=1):
        if not is_keylist_header(table):
            continue
        block = _guess_block(block_seq, idx)
        rows.extend(parse_keylist_table(table, idx, side=side, block=block))
        block_seq += 1
    return rows, docx_path
