"""跑真实 fixture 生成整合清单，并回读 docx 验证结构（表数/列数/分组行/行数）。"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

from compare_core.pipeline import run_pipeline

BASE = Path(__file__).resolve().parent.parent
FX = BASE / "runs" / "parse_real"
OUT = BASE / "runs" / "fmt_check"


def main():
    state = {
        "rows_orig_json": str(FX / "orig.json"),
        "rows_plm_json": str(FX / "plm.json"),
        "mode": "mock",
        "out_dir": str(OUT),
    }
    state = run_pipeline(state, lambda d, s: {x.key: "accept" for x in d})
    counts = {}
    for d in state["diffs"]:
        counts[d.type.value] = counts.get(d.type.value, 0) + 1
    print("diff分类:", counts, "  总diff(=输出行):", len(state["diffs"]))

    doc = Document(state["out_merged"])
    print("文档表数:", len(doc.tables))
    for ti, t in enumerate(doc.tables, 1):
        cols = len(t.columns)
        rows = len(t.rows)
        header = " | ".join(c.text for c in t.rows[0].cells)
        # 统计分组行（跨列合并 → 该行去重后单元格文本只有1种且首格非空）
        group_rows = 0
        for r in t.rows[1:]:
            texts = [c.text for c in r.cells]
            if texts[0] and len(set(texts)) == 1:
                group_rows += 1
        print(f"  表{ti}: 列={cols} 行={rows}(含表头) 分组标题行≈{group_rows}")
        print(f"    表头: {header}")
    print("INSPECT_OK")


if __name__ == "__main__":
    main()
