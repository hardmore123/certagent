"""验证审计漏洞修复：
P0-1 物理层不可读时必须降级(不得静默判 ok)
P0-2 分组标题行用合并单元格结构判定（数据行各列同值不被误杀）
P1-1 指纹加入制造商，降低碰撞
P1-2 repair 后重新审计，报告反映补录后状态
"""
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from compare_core import nodes, trace
from compare_core.differ import run_audit
from compare_core.differ.audit import _FP_FIELDS, _is_merged_title_row, _read_output
from compare_core.parser.docx_parser import parse_docx

F = ROOT.parent / "AI智能体-认证原始清单与新品清单合并"
OUT = ROOT / "runs" / "audit_fix_test"


def main():
    trace.new_run()
    OUT.mkdir(parents=True, exist_ok=True)
    state = {
        "rows_orig": parse_docx(str(F / "输入--【原始清单，人工已整理docx】4720467-85Z970QF(东芝)_安全&EMC关键件清单.docx")),
        "rows_plm": parse_docx(str(F / "输入--【PLM导出清单-新品docx】85ZXSF(01)(东芝)_安全&EMC关键件清单 - 副本.docx")),
        "out_dir": str(OUT), "mode": "mock", "suggest_source": "groups",
    }
    state.update(nodes.judge_node(state))
    state.update(nodes.suggest_node(state))
    diffs = state["diffs"]
    state["decisions"] = {d.key: "accept" for d in diffs}
    state.update(nodes.apply_node(state))

    print(f"指纹字段 = {_FP_FIELDS}")

    # 基线
    r0 = nodes.audit_node(state)["audit_report"]
    print(f"\n[基线] status={r0.status} 保留={r0.kept_count} 输出={r0.output_rows} "
          f"mismatch={r0.output_mismatch} 缺失={len(r0.missing_in_output)}")
    assert r0.status == "ok", f"基线应为 ok，实得 {r0.status}"
    assert r0.output_rows == r0.kept_count

    # === P0-1：物理层不可读 → 必须降级，不得判 ok ===
    r1 = run_audit(state["rows_orig"], state["rows_plm"], diffs,
                   state["decisions"], str(OUT / "不存在的文件.docx"))
    print(f"\n[P0-1 文档不可读] status={r1.status} output_rows={r1.output_rows}")
    hint = [x for x in r1.recommendations if "物理层未校验" in x]
    print(f"    提示: {hint[0][:60] if hint else '(无)'}")
    assert r1.status != "ok", "物理层未校验时不得判 ok（会给出虚假通过）"
    assert hint, "应明确提示物理层未校验"

    r1b = run_audit(state["rows_orig"], state["rows_plm"], diffs, state["decisions"], None)
    print(f"[P0-1 未提供路径] status={r1b.status}")
    assert r1b.status != "ok"

    # === P0-2：合并单元格结构判定 ===
    from docx import Document
    doc = Document(state["out_merged"])
    t1 = doc.tables[0]
    merged = [i for i, r in enumerate(t1.rows) if _is_merged_title_row(r)]
    print(f"\n[P0-2] 表1 识别出合并标题行 {len(merged)} 个（行号前5: {merged[:5]}）")
    assert len(merged) >= 1, "应能识别出分组标题行"

    # 构造"数据行各列同值"，旧文本判定会误杀，新结构判定不会
    probe = ROOT / "runs" / "audit_fix_test" / "probe.docx"
    d2 = Document()
    tb = d2.add_table(rows=1, cols=10)
    tb.rows[0].cells[0].text = "表头"
    cells = tb.add_row().cells      # 各列同值的数据行（未合并）
    for c in cells:
        c.text = "/"
    tb2row = tb.add_row()           # 真正的合并标题行
    m = tb2row.cells[0]
    for c in tb2row.cells[1:]:
        m = m.merge(c)
    m.text = "这是分组标题"
    d2.save(str(probe))
    fps, total, err = _read_output(str(probe))
    print(f"    探针文档：数据行计数={total}（期望 1：各列同值的数据行应被计入，合并标题行应被排除）err={err!r}")
    assert total == 1, f"期望 1，实得 {total}"

    # === P1-2：repair 后重新审计 ===
    broken = str(OUT / "broken.docx")
    shutil.copyfile(state["out_merged"], broken)
    doc3 = Document(broken)
    tt = doc3.tables[0]
    removed = None
    for i, row in enumerate(tt.rows):
        if i == 0 or _is_merged_title_row(row):
            continue
        removed = [c.text for c in row.cells][2][:20]
        row._element.getparent().remove(row._element)
        break
    doc3.save(broken)
    print(f"\n[P1-2] 人为删除输出行: {removed}")

    st = dict(state)
    st["out_merged"] = broken
    rep = run_audit(st["rows_orig"], st["rows_plm"], diffs, st["decisions"], broken)
    print(f"    删除后审计: status={rep.status} 缺失={len(rep.missing_in_output)}")
    assert rep.status == "critical" and rep.missing_in_output, "应检出输出丢行"

    st["audit_report"] = rep
    out = nodes.repair_node(st)
    r2 = out["audit_report"]
    print(f"    repair 后报告: status={r2.status} repaired={r2.repaired} "
          f"缺失={len(r2.missing_in_output)} 输出行={r2.output_rows}")
    assert r2.repaired is True
    assert r2.status == "ok", f"补录后重审应恢复 ok，实得 {r2.status}"
    assert not r2.missing_in_output, "补录后报告不应仍显示缺失"
    print(f"    首条建议: {r2.recommendations[0][:70]}")

    print("\nTEST_AUDIT_FIXES_DONE")


if __name__ == "__main__":
    main()
