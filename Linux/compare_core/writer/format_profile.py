"""格式画像：从输入 .docx（PLM 新版本=输出骨架）提取排版属性，套用到重建的整合清单。

为什么不直接复用输入表格骨架：解析阶段已把输入列语义规范化重排
（如"部件号"并入位号、表1"备注"列实为证书号、表2列序与我们不同），
输入表格与我们内部 Row schema 不是 1:1，无法直接在原表上追加行。

因此这里只提取"与列 schema 无关"的排版属性直接照搬：
- 页面：方向 / 纸张尺寸 / 页边距
- 字体：正文字体名、正文字号、表头字号（从数据行/表头行采样众数）
- 表格样式名
- 各表列宽：按 tblGrid 提取，套用时用"表头别名"把输入列宽对应到我们的输出列。

失败一律优雅降级（返回 None 或跳过），不影响清单生成。
"""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field

from docx.oxml.ns import qn

_EMU_PER_TWIP = 635  # 1 dxa(twip) = 635 EMU


@dataclass
class TableFmt:
    headers: list[str] = field(default_factory=list)
    col_widths_emu: list[int] = field(default_factory=list)  # 与 headers 一一对应


@dataclass
class FormatProfile:
    orientation: object | None = None
    page_width: object | None = None
    page_height: object | None = None
    margins: tuple | None = None  # (left, right, top, bottom) EMU
    font_name: str | None = None
    data_size_emu: int | None = None
    header_size_emu: int | None = None
    table_style: str | None = None
    tables: list[TableFmt] = field(default_factory=list)


def _grid_widths_emu(table) -> list[int]:
    """从 tblGrid 读各列宽（twips→EMU）。"""
    grid = table._tbl.tblGrid
    out = []
    for g in grid.findall(qn("w:gridCol")):
        w = g.get(qn("w:w"))
        out.append(int(w) * _EMU_PER_TWIP if w and w.isdigit() else 0)
    return out


def _sample_font(table, *, header: bool) -> tuple[Counter, Counter]:
    """采样字体名与字号（EMU）的出现次数。header=True 只看第 0 行，否则看数据行。"""
    names: Counter = Counter()
    sizes: Counter = Counter()
    rows = table.rows
    row_range = [rows[0]] if (header and rows) else rows[1:]
    for r in row_range:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    if run.font.name:
                        names[run.font.name] += 1
                    if run.font.size is not None:
                        sizes[int(run.font.size)] += 1
    return names, sizes


def extract_profile(docx_path: str) -> FormatProfile | None:
    """从输入 .docx 提取格式画像。任何异常返回 None。"""
    try:
        from docx import Document

        d = Document(docx_path)
        prof = FormatProfile()

        # 页面设置
        if d.sections:
            sec = d.sections[0]
            prof.orientation = sec.orientation
            prof.page_width = sec.page_width
            prof.page_height = sec.page_height
            prof.margins = (sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin)

        # 字体（跨全部表采样）+ 列宽（逐表）
        name_ctr: Counter = Counter()
        data_size_ctr: Counter = Counter()
        header_size_ctr: Counter = Counter()
        for t in d.tables:
            headers = [t.rows[0].cells[i].text.strip() for i in range(len(t.columns))] if t.rows else []
            prof.tables.append(TableFmt(headers=headers, col_widths_emu=_grid_widths_emu(t)))
            n1, s1 = _sample_font(t, header=True)
            n0, s0 = _sample_font(t, header=False)
            name_ctr += n0 + n1
            data_size_ctr += s0
            header_size_ctr += s1
            if prof.table_style is None and t.style is not None:
                prof.table_style = t.style.name

        if name_ctr:
            prof.font_name = name_ctr.most_common(1)[0][0]
        if data_size_ctr:
            prof.data_size_emu = data_size_ctr.most_common(1)[0][0]
        if header_size_ctr:
            prof.header_size_emu = header_size_ctr.most_common(1)[0][0]
        return prof
    except Exception:  # noqa: BLE001 - 画像提取失败不影响清单生成
        return None


# 输出表头 → 输入表头关键词（子串匹配）。用于把输入列宽对应到我们的输出列。
# 表1输入把"证书号"放在末列"备注"里，故 证书号 也匹配"备注"。
_HEADER_ALIASES = {
    "序号": ["序号"],
    "位号": ["位号"],
    "关键件名称": ["名称"],
    "名称": ["名称"],
    "型号": ["型号"],
    "规格/材料": ["规格"],
    "生产者(制造商)": ["制造商", "生产者"],
    "制造商": ["制造商", "生产者"],
    "生产企业": ["生产厂", "生产企业"],
    "生产厂": ["生产厂", "生产企业"],
    "认证标准": ["认证标准"],
    "证书号": ["证书", "备注"],
    "本次使用/备用": ["本次使用"],
    "备注": ["备注"],
}


def map_widths(out_headers: list[str], tfmt: TableFmt | None) -> list[int | None]:
    """把输入表列宽按表头别名对应到我们的输出表头。返回与 out_headers 等长的 EMU 列表（无匹配为 None）。

    贪心一对一：每个输入列宽最多被一个输出列认领，避免重复。
    """
    if not tfmt or not tfmt.headers or not tfmt.col_widths_emu:
        return [None] * len(out_headers)
    used: set[int] = set()
    result: list[int | None] = []
    for oh in out_headers:
        aliases = _HEADER_ALIASES.get(oh, [oh])
        picked: int | None = None
        for idx, ih in enumerate(tfmt.headers):
            if idx in used or idx >= len(tfmt.col_widths_emu):
                continue
            if any(a in ih for a in aliases):
                picked = tfmt.col_widths_emu[idx]
                used.add(idx)
                break
        result.append(picked)
    return result
