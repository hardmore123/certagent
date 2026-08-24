"""检查生成的整合关键件清单 docx：表数、各表行数、表1/表2 抽样数据。"""
import sys
from pathlib import Path

from docx import Document

MERGED = Path(r"c:\Users\likunyuan\Desktop\项目\result_2026-07-23\整合版关键件清单.docx")


def main():
    print("文件:", MERGED, "存在:", MERGED.exists())
    doc = Document(str(MERGED))
    print("表数量:", len(doc.tables))
    for ti, t in enumerate(doc.tables, 1):
        rows = t.rows
        cols = len(t.columns)
        header = [c.text for c in rows[0].cells]
        # 分组标题行：整行合并 → 单元格文本去重后仅1种且非空
        group = 0
        data = 0
        for r in rows[1:]:
            texts = [c.text for c in r.cells]
            if texts[0] and len(set(texts)) == 1:
                group += 1
            else:
                data += 1
        print(f"\n===== 表{ti} =====")
        print(f"  列数={cols}  总行={len(rows)}(含表头)  分组标题行={group}  数据行={data}")
        print(f"  表头: {' | '.join(header)}")
        print("  数据抽样(前5条非分组行):")
        shown = 0
        for r in rows[1:]:
            texts = [c.text for c in r.cells]
            if texts[0] and len(set(texts)) == 1:
                continue  # 跳过分组行
            print("    " + " | ".join(t[:14] for t in texts))
            shown += 1
            if shown >= 5:
                break
    print("\nCHECK_MERGED_DONE")


if __name__ == "__main__":
    main()
