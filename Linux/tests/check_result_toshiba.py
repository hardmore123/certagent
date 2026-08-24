"""检查东芝结果：整合清单行数/表结构/红色标注数，对照表条目。"""
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

BASE = Path(r"c:\Users\likunyuan\Desktop\项目")
dirs = sorted(BASE.glob("result_东芝85ZXSF_*"))
D = dirs[-1]
print("结果目录:", D)


def is_red(cell):
    for p in cell.paragraphs:
        for r in p.runs:
            if r.font.color and r.font.color.rgb == RGBColor(0xFF, 0x00, 0x00):
                return True
    return False


def main():
    m = Document(str(D / "整合版关键件清单.docx"))
    print("整合清单 表数:", len(m.tables))
    for ti, t in enumerate(m.tables, 1):
        rows = t.rows
        red = grp = data = 0
        for r in rows[1:]:
            texts = [c.text for c in r.cells]
            if texts[0] and len(set(texts)) == 1:
                grp += 1
                continue
            data += 1
            if any(is_red(c) for c in r.cells):
                red += 1
        print(f"  表{ti}: 列{len(t.columns)} 总行{len(rows)} 分组行{grp} 数据行{data} 其中红色标注{red}")

    c = Document(str(D / "变更前后对照表.docx"))
    ct = c.tables[0]
    print(f"\n对照表: {len(ct.rows)-1} 条变更项")
    for r in ct.rows[1:6]:
        print("   " + " | ".join(x.text[:22] for x in r.cells))
    print("CHECK_TOSHIBA_DONE")


if __name__ == "__main__":
    main()
