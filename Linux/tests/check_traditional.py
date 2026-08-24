"""可靠检测 docx 中的繁体字：opencc t2s 后与原字不同即为繁体。"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from opencc import OpenCC

CC = OpenCC("t2s")


def trad_chars(text: str) -> set[str]:
    """返回 text 中的繁体字集合（该字 t2s 后发生变化）。"""
    return {ch for ch in text if CC.convert(ch) != ch}


def scan_docx(path: Path):
    doc = Document(str(path))
    hits = {}
    texts = []
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                texts.append(c.text)
    for p in doc.paragraphs:
        texts.append(p.text)
    for txt in texts:
        for ch in trad_chars(txt):
            hits[ch] = hits.get(ch, 0) + txt.count(ch)
    return hits


def main():
    result_dir = Path(r"c:\Users\likunyuan\Desktop\项目") 
    # 扫描最新 result_* 目录
    dirs = sorted(result_dir.glob("result_*"))
    target = dirs[-1] if dirs else None
    print("扫描目录:", target)
    total = {}
    for f in sorted(target.glob("*.docx")):
        hits = scan_docx(f)
        print(f"\n{f.name}: {'无繁体字 ✅' if not hits else '仍有繁体字 ⚠'}")
        for ch, n in hits.items():
            print(f"    繁体 [{ch}] → 简体 [{CC.convert(ch)}]  x{n}")
            total[ch] = total.get(ch, 0) + n
    print("\n合计繁体字种类:", len(total))
    print("CHECK_TRAD_DONE")


if __name__ == "__main__":
    main()
