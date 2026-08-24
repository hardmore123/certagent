"""对比我们的输出 vs 业务期望输出，评估准确率。

期望输出(.doc, Word COM 读)：
  - 合成后清单 / 合成后变更前后对照表
我们的输出(docx, python-docx 读)：result_东芝85ZXSF_*
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

from compare_core.parser import parse_document

FOLDER = Path(r"c:\Users\likunyuan\Desktop\项目\AI智能体-认证原始清单与新品清单合并")
EXP_LIST = FOLDER / "输出--【合成后清单】85ZXSF(01)(东芝)_安全&EMC关键件清单.doc"
EXP_CMP = FOLDER / "输出--【合成后变更前后对照表】4994047-85ZXSF(01)(东芝)变更前后对照表.doc"

BASE = Path(r"c:\Users\likunyuan\Desktop\项目")
OURS = sorted(BASE.glob("result_东芝85ZXSF_*"))[-1]


def clean(s):
    return re.sub(r"\s+", "", (s or "")).replace("（", "(").replace("）", ")")


def core_name(text):
    """从'增加/变更：电线组件(...)' 或 '2.增加机内电源单元备选件' 提取核心关键件名。"""
    t = text
    if "：" in t:
        t = t.split("：", 1)[1]
    t = re.sub(r"^[0-9]+[.、]?", "", t)
    t = re.sub(r"(增加|变更|补充|备选件|备用件)", "", t)
    t = re.sub(r"\(.*?\)", "", t)
    return clean(t)


def read_expected_compare(path):
    """Word COM 读期望对照表第一列(变更项目)。"""
    import win32com.client as win32
    word = win32.Dispatch("Word.Application")
    word.Visible = False
    items = []
    try:
        doc = word.Documents.Open(str(path), ReadOnly=True)
        try:
            for ti in range(1, doc.Tables.Count + 1):
                t = doc.Tables(ti)
                cells = t.Range.Cells
                # 收集第一列文本
                from collections import defaultdict
                by_row = defaultdict(dict)
                for i in range(1, cells.Count + 1):
                    c = cells.Item(i)
                    by_row[int(c.RowIndex)][int(c.ColumnIndex)] = c.Range.Text.replace("\r", "").replace("\x07", "").strip()
                for ri in sorted(by_row):
                    col1 = by_row[ri].get(1, "")
                    if col1 and col1 not in ("变更项目",):
                        items.append(col1)
        finally:
            doc.Close(False)
    finally:
        word.Quit()
    return items


def main():
    # 1) 期望合成清单行数
    print("解析期望合成清单...")
    exp_rows = parse_document(str(EXP_LIST))
    exp_t1 = [r for r in exp_rows if r.table == 1]
    exp_t2 = [r for r in exp_rows if r.table == 2]
    print(f"  期望合成清单: 总{len(exp_rows)} (表1 {len(exp_t1)} + 表2 {len(exp_t2)})")

    # 我们的合成清单行数
    m = Document(str(OURS / "整合版关键件清单.docx"))
    our_counts = []
    for t in m.tables:
        data = sum(1 for r in t.rows[1:] if not (r.cells[0].text and len({c.text for c in r.cells}) == 1))
        our_counts.append(data)
    print(f"  我们合成清单: 表1数据行 {our_counts[0]} + 表2数据行 {our_counts[1] if len(our_counts)>1 else 0}")

    # 2) 对照表对比
    print("\n解析期望对照表...")
    exp_items = read_expected_compare(EXP_CMP)
    print(f"  期望对照表变更项: {len(exp_items)}")
    for it in exp_items:
        print("    [期望]", it[:40])

    c = Document(str(OURS / "变更前后对照表.docx"))
    our_items = [r.cells[0].text for r in c.tables[0].rows[1:]]
    print(f"\n  我们对照表变更项: {len(our_items)}")

    exp_cores = [core_name(x) for x in exp_items]
    our_cores = [core_name(x) for x in our_items]

    matched = 0
    for oc in our_cores:
        if oc and any(oc in ec or ec in oc for ec in exp_cores if ec):
            matched += 1
    covered = 0
    for ec in exp_cores:
        if ec and any(ec in oc or oc in ec for oc in our_cores if oc):
            covered += 1
    prec = matched / max(len(our_cores), 1)
    rec = covered / max(len(exp_cores), 1)
    print(f"\n对照表匹配: 我们命中期望 精确率≈{prec:.0%} ({matched}/{len(our_cores)})  "
          f"覆盖期望 召回率≈{rec:.0%} ({covered}/{len(exp_cores)})")
    print("COMPARE_ACC_DONE")


if __name__ == "__main__":
    main()
