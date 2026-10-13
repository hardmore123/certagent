"""差异汇总 Word 文档 writer。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from .models import VerifyDiff


def _short(value: str, limit: int = 1800) -> str:
    value = value or ""
    return value if len(value) <= limit else value[:limit] + "…"


def _suggestion(diff: VerifyDiff) -> str:
    if diff.diff_type == "block_missing_in_draft":
        return "请确认该子块是否适用于当前报告草稿；确认后再决定是否补充"
    if diff.diff_type == "block_extra_in_draft":
        return "请确认该子块是否为模板差异或应补充认证资料依据"
    if diff.diff_type == "missing_in_draft":
        return "请确认是否补充到报告草稿"
    if diff.diff_type == "extra_in_draft":
        return "请确认是否从报告草稿删除或补充认证资料依据"
    return "请按认证资料核实并修改报告草稿"


def write_summary(
    diffs: list[VerifyDiff],
    out_path: str,
    *,
    cert_name: str = "",
    draft_name: str = "",
    warnings: list[str] | None = None,
) -> str:
    document = Document()
    document.add_heading("报告草稿准确性核对——差异汇总", level=1)
    p = document.add_paragraph()
    p.add_run("认证资料（CERT，应然基准）：").bold = True
    p.add_run(cert_name)
    p = document.add_paragraph()
    p.add_run("报告草稿（DRAFT，被核对对象）：").bold = True
    p.add_run(draft_name)
    document.add_paragraph(
        f"共发现 {len(diffs)} 条差异。默认采用严格逐字核对；专业路由状态：discipline_pending。"
    )
    if warnings:
        document.add_heading("人工复核提示", level=2)
        for warning in warnings:
            document.add_paragraph(warning, style="List Bullet")

    if not diffs:
        document.add_paragraph("未发现字段差异。")
    else:
        headers = [
            "序号", "子块", "项目", "字段", "认证资料 CERT（应然）",
            "报告草稿 DRAFT（实然）", "差异类型", "匹配状态", "匹配置信度",
            "匹配说明", "DRAFT 批注状态", "建议",
        ]
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = header
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)
        for index, diff in enumerate(diffs, start=1):
            row = table.add_row().cells
            comment_status = "已写入" if diff.comment_written else (
                diff.comment_error or "无 DRAFT 锚点（仅汇总）"
            )
            values = [
                str(index), diff.block_type, _short(diff.item_name, 300), diff.field,
                _short(diff.cert_value), _short(diff.draft_value), diff.diff_type,
                diff.match_status, f"{diff.match_confidence:.2f}",
                _short(diff.match_reason, 500), _short(comment_status, 300),
                _suggestion(diff),
            ]
            for cell, value in zip(row, values):
                cell.text = value
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(8)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(Path(out_path).resolve()))
    return out_path
