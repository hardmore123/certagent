"""把差异批注写回 DRAFT 原始文档。"""
from __future__ import annotations

from pathlib import Path

from docx import Document

from .models import VerifyDiff


def _comment_text(diff: VerifyDiff) -> str:
    if diff.diff_type == "missing_in_draft":
        return f"[报告草稿差异] 认证资料存在而报告草稿缺失。认证资料值：{diff.cert_value}"
    if diff.diff_type == "extra_in_draft":
        return f"[报告草稿差异] 报告草稿存在而认证资料未找到。报告草稿值：{diff.draft_value}"
    return (
        f"[报告草稿差异] 字段：{diff.field}\n"
        f"认证资料（CERT，应然）：{diff.cert_value}\n"
        f"报告草稿（DRAFT，实然）：{diff.draft_value}"
    )


def _anchor_runs(document, diff: VerifyDiff):
    anchor = diff.draft_anchor
    if anchor is None:
        return None
    table = document.tables[anchor.table_idx - 1]
    row = table.rows[anchor.row_idx - 1]
    cell = row.cells[anchor.cell_idx]
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    runs = list(paragraph.runs)
    if not runs:
        runs = [paragraph.add_run(" ")]
    return runs


def write_comments(
    draft_docx_path: str,
    diffs: list[VerifyDiff],
    out_path: str,
    *,
    author: str = "ReportVerify",
    initials: str = "RV",
) -> dict[str, object]:
    """输出批注版 DRAFT；没有 DRAFT 锚点的差异交由汇总文档兜底。"""
    document = Document(str(Path(draft_docx_path).resolve()))
    failed: list[str] = []
    written = 0
    for diff in diffs:
        if diff.draft_anchor is None:
            diff.comment_error = "DRAFT 中不存在可写入的对应锚点"
            continue
        try:
            runs = _anchor_runs(document, diff)
            if not runs:
                raise ValueError("无法取得批注段落")
            document.add_comment(runs, _comment_text(diff), author=author, initials=initials)
            diff.comment_written = True
            written += 1
        except Exception as exc:  # noqa: BLE001 - 单条失败不能吞掉汇总
            diff.comment_error = f"批注锚点失败：{type(exc).__name__}: {exc}"
            failed.append(diff.key)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(Path(out_path).resolve()))
    return {"written": written, "failed": failed, "total": len(diffs)}
