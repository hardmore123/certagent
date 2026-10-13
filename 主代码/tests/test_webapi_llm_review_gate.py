"""任务 3.4：LLM 状态与人工确认/下载门禁回归。

所有 provider 行为均由本地返回值或异常模拟；测试不打开 socket、不访问真实
provider，也不使用真实凭据。重点验证 LLM 只是 review 附属 warning，不能改变
确定性任务生命周期、人工决策覆盖范围或最终下载门禁。
"""
from __future__ import annotations

import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from report_verify.models import Anchor, VerifyDiff, to_dict
from report_verify.llm.config import LLMSettings
from webapi.app import app, download_result
from webapi.jobs import Job, store
import webapi.app as web_app


_PLACEHOLDER_API_KEY = "PLACEHOLDER-API-KEY-DO-NOT-USE"
_SENSITIVE_PLACEHOLDER = "PLACEHOLDER-PROMPT-RESPONSE-CERT-ORIGINAL"


def _enabled_settings() -> LLMSettings:
    return LLMSettings(
        mode="candidate",
        base_url="https://approved.example/v1",
        model="local-stub-model",
        api_key=_PLACEHOLDER_API_KEY,
        allowed_hosts=frozenset({"approved.example"}),
    )


def _job(root: str, name: str) -> Job:
    job = Job(
        job_id=f"review-gate-{name}",
        cert_path=str(Path(root) / f"{name}-cert.docx"),
        draft_path=str(Path(root) / f"{name}-draft.docx"),
        out_dir=str(Path(root) / f"{name}-run"),
    )
    # The lifecycle tests use an in-memory Job; no process-wide job directory is
    # needed and no provider secret is persisted.
    job.save = Mock()  # type: ignore[method-assign]
    return job


def _deterministic_trace(root: str) -> dict:
    return {
        "result": {
            "diffs": [
                {
                    "key": "deterministic-diff-1",
                    "discipline": "安全",
                    "discipline_status": "discipline_pending",
                }
            ]
        },
        "rules": {"version": "rules-local"},
    }


def _review_trace(root: Path) -> Path:
    draft = root / "draft.docx"
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "项目一"
    table.cell(1, 1).text = "旧值一"
    table.cell(2, 0).text = "项目二"
    table.cell(2, 1).text = "旧值二"
    document.save(draft)

    diffs = [
        VerifyDiff(
            key="diff-1",
            block_type="safety_list",
            item_name="项目一",
            diff_type="field_mismatch",
            field="值",
            cert_value="新值一",
            draft_value="旧值一",
            draft_anchor=Anchor(1, 2, 1, "旧值一"),
            discipline="安全",
            discipline_status="discipline_pending",
            rule_version="rules-local",
        ),
        VerifyDiff(
            key="diff-2",
            block_type="safety_list",
            item_name="项目二",
            diff_type="field_mismatch",
            field="值",
            cert_value="新值二",
            draft_value="旧值二",
            draft_anchor=Anchor(1, 3, 1, "旧值二"),
            discipline="安全",
            discipline_status="discipline_pending",
            rule_version="rules-local",
        ),
    ]
    trace = {
        "cert_path": str(root / "cert.docx"),
        "draft_path": str(draft),
        "draft_docx": str(draft),
        "outputs": {"annotated_draft": str(root / "批注版报告草稿.docx")},
        "result": {"diffs": to_dict(diffs), "warnings": []},
    }
    trace_path = root / "verification.json"
    trace_path.write_text(json.dumps(trace, ensure_ascii=False), encoding="utf-8")
    return trace_path


class LLMReviewLifecycleTests(unittest.IsolatedAsyncioTestCase):
    async def test_all_llm_outcomes_remain_review_and_only_emit_safe_warning(self):
        outcomes = (
            ("success", _enabled_settings(), {"status": "success", "capabilities": ["candidate"]}),
            (
                "fallback",
                _enabled_settings(),
                {
                    "status": "fallback",
                    "error_code": "llm_http_error",
                    "warnings": [{"code": "llm_http_error", "message": _SENSITIVE_PLACEHOLDER}],
                    "capabilities": ["candidate", "prompt"],
                },
            ),
            (
                "skipped_policy",
                _enabled_settings(),
                {
                    "status": "skipped_policy",
                    "error_code": "llm_policy_denied",
                    "warnings": [{"code": "llm_policy_denied", "message": _SENSITIVE_PLACEHOLDER}],
                },
            ),
            (
                "invalid",
                _enabled_settings(),
                {
                    "status": "invalid",
                    "error_code": "llm_schema_invalid",
                    "warnings": [{"code": "llm_schema_invalid", "message": _SENSITIVE_PLACEHOLDER}],
                },
            ),
            (
                "misconfigured",
                LLMSettings(mode="candidate"),
                None,
            ),
            ("disabled", LLMSettings(mode="off"), None),
        )

        for expected_status, runtime_settings, evidence in outcomes:
            with self.subTest(status=expected_status), tempfile.TemporaryDirectory() as root:
                orchestrator = Mock(return_value=evidence)
                with (
                    patch.object(web_app, "semaphore", asyncio.Semaphore(1)),
                    patch.object(web_app, "provider_semaphore", asyncio.Semaphore(1)),
                    patch.object(web_app, "llm_settings", runtime_settings),
                    patch.object(web_app, "run_verification", return_value=_deterministic_trace(root)),
                    patch.object(web_app, "_run_llm_orchestrator", orchestrator),
                ):
                    job = _job(root, expected_status)
                    await web_app._execute(job)

                self.assertEqual(job.status, "review")
                self.assertEqual(job.error, "")
                self.assertEqual(job.diff_count, 1)
                self.assertFalse(job.summary()["download_ready"])
                self.assertEqual(job.decisions, {})
                self.assertNotIn(_SENSITIVE_PLACEHOLDER, json.dumps(job.summary(), ensure_ascii=False))
                self.assertNotIn("prompt", json.dumps(job.summary(), ensure_ascii=False).lower())
                self.assertEqual(job.llm_status, expected_status)

                if expected_status == "disabled":
                    orchestrator.assert_not_called()
                    self.assertEqual(job.llm_warning_count, 0)
                elif expected_status == "misconfigured":
                    # Missing server credentials are rejected before the worker
                    # can enter the provider path.
                    orchestrator.assert_not_called()
                    self.assertEqual(job.llm_warning_count, 1)
                    self.assertEqual(job.llm_last_error_code, "llm_misconfigured")
                else:
                    orchestrator.assert_called_once()
                    self.assertGreaterEqual(job.llm_warning_count, int(expected_status != "success"))
                    if expected_status != "success":
                        self.assertNotEqual(job.llm_last_error_code, _SENSITIVE_PLACEHOLDER)

    async def test_provider_exception_is_a_review_warning_not_deterministic_failure(self):
        with tempfile.TemporaryDirectory() as root:
            with (
                patch.object(web_app, "semaphore", asyncio.Semaphore(1)),
                patch.object(web_app, "provider_semaphore", asyncio.Semaphore(1)),
                patch.object(web_app, "llm_settings", _enabled_settings()),
                patch.object(web_app, "run_verification", return_value=_deterministic_trace(root)),
                patch.object(
                    web_app,
                    "_run_llm_orchestrator",
                    side_effect=RuntimeError(_SENSITIVE_PLACEHOLDER),
                ),
            ):
                job = _job(root, "provider-error")
                await web_app._execute(job)

            self.assertEqual(job.status, "review")
            self.assertEqual(job.llm_status, "fallback")
            self.assertEqual(job.llm_warning_count, 1)
            self.assertEqual(job.llm_last_error_code, "llm_http_error")
            self.assertEqual(job.error, "")
            serialized = json.dumps(job.summary(), ensure_ascii=False)
            self.assertNotIn(_SENSITIVE_PLACEHOLDER, serialized)
            self.assertNotIn("RuntimeError", serialized)


class HumanReviewAndDownloadGateTests(unittest.TestCase):
    def test_partial_confirmation_is_not_downloadable_and_full_finalize_is_required(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            trace_path = _review_trace(root)
            job = Job(
                job_id="review-gate-api",
                status="review",
                stage="等待人工确认（LLM 已回退）",
                progress=0.7,
                out_dir=str(root),
                trace_path=str(trace_path),
                diff_count=2,
                llm_mode="candidate",
                llm_status="fallback",
                llm_warning_count=1,
                llm_last_error_code="llm_http_error",
            )
            job.save = Mock()  # type: ignore[method-assign]
            store.jobs[job.job_id] = job
            client = TestClient(app)
            try:
                partial = client.post(
                    f"/api/jobs/{job.job_id}/decisions",
                    json={"decisions": {"diff-1": {"verdict": "accept", "note": "人工确认"}}},
                    headers={"X-Reviewer": "local-reviewer"},
                )
                self.assertEqual(partial.status_code, 200)
                partial_payload = partial.json()
                self.assertEqual(partial_payload["status"], "review")
                self.assertEqual(partial_payload["missing_keys"], ["diff-2"])
                self.assertFalse(partial_payload["download_ready"])
                self.assertEqual(job.decisions_count, 1)
                self.assertEqual(job.llm_status, "fallback")

                blocked = client.get(f"/api/jobs/{job.job_id}/result/draft")
                self.assertEqual(blocked.status_code, 409)
                self.assertFalse((root / "批注版报告草稿_最终.docx").exists())

                completed = client.post(
                    f"/api/jobs/{job.job_id}/decisions",
                    json={"decisions": {"diff-2": {"verdict": "ignore", "note": "人工忽略"}}},
                    headers={"X-Reviewer": "local-reviewer"},
                )
                self.assertEqual(completed.status_code, 200)
                completed_payload = completed.json()
                self.assertEqual(completed_payload["status"], "done")
                self.assertTrue(completed_payload["download_ready"])
                self.assertEqual(job.decisions_count, 2)
                self.assertEqual(job.llm_status, "fallback")

                downloaded_draft = client.get(f"/api/jobs/{job.job_id}/result/draft")
                downloaded_summary = client.get(f"/api/jobs/{job.job_id}/result/summary")
                self.assertEqual(downloaded_draft.status_code, 200)
                self.assertEqual(downloaded_summary.status_code, 200)
                self.assertTrue((root / "批注版报告草稿_最终.docx").exists())
                self.assertTrue((root / "差异汇总_最终.docx").exists())

                review = json.loads((root / "review.json").read_text(encoding="utf-8"))
                self.assertEqual(review["review"]["status"], "completed")
                self.assertEqual(review["diffs"][0]["discipline_status"], "discipline_pending")
                self.assertEqual(review["diffs"][1]["discipline_status"], "discipline_pending")
                self.assertEqual(review["diffs"][0]["reviewed_by"], "local-reviewer")
            finally:
                store.jobs.pop(job.job_id, None)

    def test_llm_state_alone_cannot_call_download_result(self):
        job = Job(
            job_id="review-gate-no-finalize",
            status="review",
            out_dir=".",
            llm_mode="candidate_explain",
            llm_status="success",
            llm_suggestion_count=3,
        )
        store.jobs[job.job_id] = job
        try:
            with self.assertRaises(HTTPException) as raised:
                asyncio.run(download_result(job.job_id, "draft"))
            self.assertEqual(raised.exception.status_code, 409)
            self.assertFalse(job.summary()["download_ready"])
        finally:
            store.jobs.pop(job.job_id, None)


if __name__ == "__main__":
    unittest.main()
