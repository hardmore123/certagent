"""任务 4.4：候选守恒、方向、规则版本与人工下载门禁回归。

所有 provider 行为都由内存 client stub 提供；测试只使用占位数据，不开 socket、
不访问真实 provider，也不读取、复制或写入真实凭据。这里把任务 4.4 的不变量
串成最小的回归场景，补充各阶段单元测试之间的组合边界。
"""
from __future__ import annotations

import asyncio
import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException

import report_verify.llm.recheck as recheck_module
from report_verify.llm.client import (
    ERROR_HTTP_ERROR,
    LLMCallMetadata,
    LLMClientError,
    LLMRawResult,
)
from report_verify.llm.config import LLMSettings
from report_verify.llm.discipline import (
    CANDIDATE_STATUS_PENDING as DISCIPLINE_PENDING,
    PublishedDisciplineRules,
    STATUS_SUCCESS as DISCIPLINE_SUCCESS,
    generate_discipline_suggestions,
)
from report_verify.llm.orchestrator import (
    STATUS_FALLBACK,
    STATUS_INVALID,
    STATUS_SKIPPED_POLICY,
    STATUS_SUCCESS,
    run_candidate_orchestrator,
)
from report_verify.llm.policy import ERROR_POLICY_DENIED, ServerLLMPolicy
from report_verify.llm.recheck import (
    CANDIDATE_STATUS_ACCEPTED_FOR_RECHECK,
    accept_candidate_for_recheck,
)
from report_verify.models import Anchor, DocumentRow, LocatedBlock, VerifyDiff, VerifyResult, to_dict
from report_verify.review import finalize_review
from webapi.app import download_result
from webapi.jobs import Job, store


_SCHEMA = "reportverify-llm-v1"
_RULE_VERSION = "rules-4-4"
_POLICY_VERSION = "policy-4-4"
_BLOCK = "safety_list"
_PLACEHOLDER_API_KEY = "PLACEHOLDER-API-KEY-DO-NOT-USE"


class _MemoryClient:
    """注入 orchestrator 的本地 provider stub；不会创建网络连接。"""

    def __init__(self, content: str = "", error: LLMClientError | None = None) -> None:
        self.content = content
        self.error = error
        self.calls: list[dict[str, object]] = []

    def complete_raw(self, messages, *, schema_version, request_id):
        self.calls.append({
            "messages": messages,
            "schema_version": schema_version,
            "request_id": request_id,
        })
        if self.error is not None:
            raise self.error
        return LLMRawResult(
            content=self.content,
            metadata=LLMCallMetadata(
                provider_host="stub.local",
                model="local-stub-model",
                schema_version=schema_version,
                attempts=1,
                elapsed_ms=1,
                transport="json",
                status="success",
                request_sha256="request-hash",
                response_sha256="response-hash",
            ),
        )


def _settings(mode: str = "candidate") -> LLMSettings:
    return LLMSettings(
        mode=mode,  # type: ignore[arg-type]
        base_url="http://127.0.0.1:8765/v1",
        model="local-stub-model",
        api_key=_PLACEHOLDER_API_KEY,
        allowed_hosts=frozenset({"127.0.0.1"}),
        allow_plaintext=True,
        policy_version=_POLICY_VERSION,
        response_schema_version=_SCHEMA,
    )


def _row(name: str, model: str, *, row_idx: int, group: str = "G1") -> DocumentRow:
    return DocumentRow(
        block_type=_BLOCK,
        row_idx=row_idx,
        values={"关键件名称": name, "型号规格": model},
        source_table=1,
        group=group,
    )


def _display_row(row: DocumentRow) -> str:
    payload = {
        str(key): str(value or "")
        for key, value in row.values.items()
        if not key.endswith("__2")
    }
    if row.group and "分组" not in payload:
        payload["分组"] = row.group
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"))


def _residual_result(
    cert_rows: list[DocumentRow], draft_rows: list[DocumentRow],
) -> VerifyResult:
    diffs: list[VerifyDiff] = []
    for row in cert_rows:
        diffs.append(VerifyDiff(
            key=f"cert-{row.row_idx}",
            block_type=_BLOCK,
            item_name=row.item_name,
            diff_type="missing_in_draft",
            field="整行",
            cert_value=_display_row(row),
            cert_location=f"表1第{row.row_idx}行",
            match_status="unmatched",
            rule_version=_RULE_VERSION,
        ))
    for row in draft_rows:
        diffs.append(VerifyDiff(
            key=f"draft-{row.row_idx}",
            block_type=_BLOCK,
            item_name=row.item_name,
            diff_type="extra_in_draft",
            field="整行",
            draft_value=_display_row(row),
            draft_anchor=Anchor(1, row.row_idx, 1),
            match_status="unmatched",
            rule_version=_RULE_VERSION,
        ))
    return VerifyResult(
        cert_blocks=[LocatedBlock(block_type=_BLOCK, table_idx=1, rows=list(cert_rows))],
        draft_blocks=[LocatedBlock(block_type=_BLOCK, table_idx=1, rows=list(draft_rows))],
        diffs=diffs,
        rule_version=_RULE_VERSION,
    )


def _candidate_response() -> str:
    return json.dumps({
        "schema_version": _SCHEMA,
        "pairs": [
            {"cert_index": 0, "draft_index": 0, "confidence": 0.95, "reason": "同表同组"},
            {"cert_index": 1, "draft_index": 1, "confidence": 0.90, "reason": "同表同组"},
        ],
    }, ensure_ascii=False)


class CandidateConservationPropertyTests(unittest.TestCase):
    def _result(self) -> VerifyResult:
        return _residual_result(
            [_row("CERT-A", "A-100", row_idx=2), _row("CERT-B", "B-200", row_idx=3)],
            [_row("DRAFT-A", "A-101", row_idx=5), _row("DRAFT-B", "B-201", row_idx=6)],
        )

    def test_success_preserves_rows_diffs_direction_and_versions(self):
        result = self._result()
        before = copy.deepcopy(result)
        out = run_candidate_orchestrator(
            result,
            _MemoryClient(_candidate_response()),
            _settings(),
            trace_id="trace-4-4-success",
            threshold=0.6,
        )

        self.assertEqual(out.status, STATUS_SUCCESS)
        self.assertEqual(len(out.candidates), 2)
        self.assertEqual(result, before)
        self.assertEqual(out.metadata["rule_version"], _RULE_VERSION)
        self.assertEqual(out.metadata["policy_version"], _POLICY_VERSION)

        cert_by_row = {row.row_idx: row for row in result.cert_blocks[0].rows}
        draft_by_row = {row.row_idx: row for row in result.draft_blocks[0].rows}
        cert_refs = set()
        draft_refs = set()
        for candidate in out.candidates:
            cert_row_idx = candidate.cert_ref["row_idx"]
            draft_row_idx = candidate.draft_ref["row_idx"]
            self.assertIn(cert_row_idx, cert_by_row)
            self.assertIn(draft_row_idx, draft_by_row)
            cert_row = cert_by_row[cert_row_idx]
            draft_row = draft_by_row[draft_row_idx]
            self.assertEqual(candidate.cert_ref["block_type"], _BLOCK)
            self.assertEqual(candidate.draft_ref["block_type"], _BLOCK)
            self.assertEqual(cert_row.source_table, draft_row.source_table)
            self.assertEqual(cert_row.group, draft_row.group)
            # CERT 是应然侧引用，DRAFT 是实然侧引用；两侧不可交换。
            self.assertNotEqual(cert_row_idx, draft_row_idx)
            cert_refs.add(cert_row_idx)
            draft_refs.add(draft_row_idx)
            self.assertEqual(candidate.status, "pending")
            self.assertEqual(candidate.policy_version, _POLICY_VERSION)
            self.assertEqual(candidate.schema_version, _SCHEMA)

        self.assertEqual(cert_refs, {2, 3})
        self.assertEqual(draft_refs, {5, 6})
        self.assertEqual(len(cert_refs), len(out.candidates))
        self.assertEqual(len(draft_refs), len(out.candidates))

    def test_failure_and_policy_denial_preserve_complete_deterministic_evidence(self):
        failure_result = self._result()
        failure_before = copy.deepcopy(failure_result)
        failure = run_candidate_orchestrator(
            failure_result,
            _MemoryClient(error=LLMClientError(ERROR_HTTP_ERROR, "local stub failure")),
            _settings(),
            trace_id="trace-4-4-failure",
        )
        self.assertEqual(failure.status, STATUS_FALLBACK)
        self.assertEqual(failure_result, failure_before)
        self.assertEqual(failure.metadata["rule_version"], _RULE_VERSION)

        denied_result = self._result()
        denied_before = copy.deepcopy(denied_result)
        denied_client = _MemoryClient(_candidate_response())
        denied_policy = ServerLLMPolicy(
            mode="candidate",
            allowed_hosts=frozenset({"stub.local"}),
            allow_plaintext=True,
            max_rows=1,  # 残余池有四行，策略在网络调用前拒绝。
            policy_version=_POLICY_VERSION,
        )
        denied = run_candidate_orchestrator(
            denied_result,
            denied_client,
            denied_policy,
            trace_id="trace-4-4-policy-denied",
        )
        self.assertEqual(denied.status, STATUS_SKIPPED_POLICY)
        self.assertEqual(denied.error_code, ERROR_POLICY_DENIED)
        self.assertEqual(denied_result, denied_before)
        self.assertEqual(denied_client.calls, [])

    def test_model_cannot_rewrite_deterministic_rule_version(self):
        result = self._result()
        before = copy.deepcopy(result)
        response = json.dumps({
            "schema_version": _SCHEMA,
            "rule_version": "attacker-rule-version",
            "pairs": [],
        }, ensure_ascii=False)
        out = run_candidate_orchestrator(
            result, _MemoryClient(response), _settings(), trace_id="trace-4-4-rule-version"
        )

        self.assertEqual(out.status, STATUS_INVALID)
        self.assertEqual(result, before)
        self.assertEqual(out.metadata["rule_version"], _RULE_VERSION)
        self.assertNotIn("attacker-rule-version", json.dumps(out.to_dict(), ensure_ascii=False))

    def test_accepted_candidate_only_runs_deterministic_recheck(self):
        result = _residual_result(
            [_row("CERT-A", "A-100", row_idx=2)],
            [_row("DRAFT-A", "A-200", row_idx=5)],
        )
        candidate_result = run_candidate_orchestrator(
            result,
            _MemoryClient(json.dumps({
                "schema_version": _SCHEMA,
                "pairs": [{
                    "cert_index": 0,
                    "draft_index": 0,
                    "confidence": 0.9,
                    "reason": "同表同组",
                }],
            }, ensure_ascii=False)),
            _settings(),
            trace_id="trace-4-4-accept",
        )
        candidate = candidate_result.candidates[0]
        before = copy.deepcopy(result)

        with patch.object(
            recheck_module, "same_value", wraps=recheck_module.same_value
        ) as deterministic_same_value:
            outcome = accept_candidate_for_recheck(
                candidate.candidate_key,
                [candidate],
                result,
                rules={"version": _RULE_VERSION, "enabled": []},
            )

        self.assertEqual(outcome.status, CANDIDATE_STATUS_ACCEPTED_FOR_RECHECK)
        self.assertEqual(outcome.acceptance_status, CANDIDATE_STATUS_ACCEPTED_FOR_RECHECK)
        self.assertTrue(outcome.recheck_requested)
        self.assertEqual(outcome.recheck_result["status"], "success")
        self.assertEqual(outcome.recheck_result["rule_version"], _RULE_VERSION)
        self.assertTrue(deterministic_same_value.called)
        self.assertEqual(candidate.status, "pending")
        self.assertEqual(result, before)
        self.assertTrue(outcome.manual_confirmation_required)
        self.assertFalse(outcome.download_ready)


class DisciplineAndDownloadBoundaryTests(unittest.TestCase):
    def test_discipline_suggestion_stays_pending_and_does_not_write_final_field(self):
        diff = VerifyDiff(
            key="diff-4-4",
            block_type=_BLOCK,
            item_name="电源",
            diff_type="field_mismatch",
            field="型号规格",
            cert_value="A-100",
            draft_value="A-200",
            discipline=None,
            discipline_status="discipline_pending",
            rule_version=_RULE_VERSION,
        )
        result = VerifyResult(diffs=[diff], rule_version=_RULE_VERSION)
        before = copy.deepcopy(result)
        rules = PublishedDisciplineRules(
            allowed_disciplines=frozenset({"安全"}),
            version="discipline-rules-4-4",
            confirmed=True,
        )
        response = json.dumps({
            "schema_version": _SCHEMA,
            "candidates": [{
                "diff_key": "diff-4-4",
                "discipline": "安全",
                "confidence": 0.9,
                "reason": "本地 stub 专业候选",
            }],
        }, ensure_ascii=False)
        out = generate_discipline_suggestions(
            result,
            rules,
            settings=_settings(mode="explain"),
            caller=lambda _metadata, _rules: response,
            trace_id="trace-4-4-discipline",
        )

        self.assertEqual(out.status, DISCIPLINE_SUCCESS)
        self.assertEqual(len(out.suggestions), 1)
        suggestion = out.suggestions[0]
        self.assertEqual(suggestion.status, DISCIPLINE_PENDING)
        self.assertEqual(suggestion.rule_version, "discipline-rules-4-4")
        self.assertEqual(result, before)
        self.assertIsNone(result.diffs[0].discipline)
        self.assertEqual(result.diffs[0].discipline_status, "discipline_pending")

    def test_llm_states_cannot_download_until_full_finalize(self):
        statuses = ("disabled", "skipped_policy", "success", "fallback", "invalid", "misconfigured")
        for index, status in enumerate(statuses):
            job = Job(
                job_id=f"task44-gate-{index}",
                status="review",
                out_dir=".",
                llm_status=status,
            )
            store.jobs[job.job_id] = job
            try:
                self.assertFalse(job.summary()["download_ready"])
                with self.assertRaises(HTTPException) as raised:
                    asyncio.run(download_result(job.job_id, "draft"))
                self.assertEqual(raised.exception.status_code, 409)
            finally:
                store.jobs.pop(job.job_id, None)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            draft_path = root / "draft.docx"
            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "字段"
            table.cell(0, 1).text = "值"
            table.cell(1, 0).text = "项目一"
            table.cell(1, 1).text = "旧值"
            document.save(draft_path)

            diff = VerifyDiff(
                key="diff-4-4-download",
                block_type=_BLOCK,
                item_name="项目一",
                diff_type="field_mismatch",
                field="值",
                cert_value="新值",
                draft_value="旧值",
                draft_anchor=Anchor(1, 2, 1, "旧值"),
                rule_version=_RULE_VERSION,
            )
            trace = {
                "cert_path": str(root / "cert.docx"),
                "draft_path": str(draft_path),
                "draft_docx": str(draft_path),
                "outputs": {"annotated_draft": str(root / "annotated.docx")},
                "result": {"diffs": to_dict([diff]), "warnings": []},
            }
            job = Job(
                job_id="task44-gate-finalize",
                status="review",
                out_dir=str(root),
                llm_status="success",
                llm_suggestion_count=1,
            )
            store.jobs[job.job_id] = job
            try:
                with self.assertRaises(HTTPException):
                    asyncio.run(download_result(job.job_id, "draft"))

                review = finalize_review(
                    trace,
                    {diff.key: {"verdict": "accept", "note": "人工确认"}},
                    reviewer="local-reviewer",
                    feedback_path=root / "feedback.jsonl",
                )
                self.assertEqual(review["status"], "completed")
                self.assertEqual(job.status, "review")
                self.assertFalse(job.summary()["download_ready"])

                # 与 Web 路由一致：只有 finalize_review 成功后才推进到 done。
                job.status = "done"
                self.assertTrue(job.summary()["download_ready"])
                draft_response = asyncio.run(download_result(job.job_id, "draft"))
                summary_response = asyncio.run(download_result(job.job_id, "summary"))
                self.assertIsNotNone(draft_response)
                self.assertIsNotNone(summary_response)
                self.assertTrue((root / "批注版报告草稿_最终.docx").exists())
                self.assertTrue((root / "差异汇总_最终.docx").exists())
            finally:
                store.jobs.pop(job.job_id, None)


if __name__ == "__main__":
    unittest.main()
