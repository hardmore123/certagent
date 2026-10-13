"""任务 3.5：反馈记忆只接受人工 verdict 的回归测试。

所有数据均为本地占位值；不访问 provider，也不使用真实凭据或真实文档。
"""
from __future__ import annotations

import asyncio
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from report_verify.llm.evidence import write_evidence_for_verification
from report_verify.models import Anchor, VerifyDiff, to_dict
from report_verify.review import (
    finalize_review,
    load_feedback,
    memory_hints,
    normalize_manual_decisions,
    record_feedback,
)
from webapi.app import app, download_result
from webapi.jobs import Job, store


_CERT_PLACEHOLDER = "CERT-ORIGINAL-PLACEHOLDER"
_DRAFT_PLACEHOLDER = "DRAFT-ORIGINAL-PLACEHOLDER"
_AI_REASON = "AI-REASON-MUST-NOT-BECOME-MEMORY"
_AI_CANDIDATE = "cand-local-only"


def _diff() -> VerifyDiff:
    return VerifyDiff(
        key="diff-local-1",
        block_type="safety_list",
        item_name="电源",
        diff_type="field_mismatch",
        field="型号",
        cert_value=_CERT_PLACEHOLDER,
        draft_value=_DRAFT_PLACEHOLDER,
        draft_anchor=Anchor(1, 2, 0, "旧值"),
        rule_version="normalized-v2",
    )


def _trace(root: Path, *, include_llm: bool = False) -> dict:
    draft = root / "draft.docx"
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "字段"
    table.cell(1, 0).text = "旧值"
    document.save(draft)
    diff = _diff()
    trace = {
        "cert_path": "local-cert.docx",
        "draft_path": "local-draft.docx",
        "draft_docx": str(draft),
        "outputs": {"annotated_draft": str(root / "initial.docx")},
        "result": {"diffs": [to_dict(diff)], "warnings": []},
    }
    if include_llm:
        trace["llm"] = {
            "mode": "candidate_explain",
            "status": "success",
            "candidates": [
                {
                    "candidate_key": _AI_CANDIDATE,
                    "reason": _AI_REASON,
                    "status": "pending",
                }
            ],
            "explanations": [{"diff_key": diff.key, "explanation": _AI_REASON}],
            "discipline_suggestions": [
                {"diff_key": diff.key, "candidate_discipline": "安全", "status": "pending"}
            ],
            "calls": [{"provider_host": "stub.local", "status": "success"}],
        }
    return trace


class FeedbackMemoryBoundaryTests(unittest.TestCase):
    def test_llm_evidence_without_human_verdict_creates_no_memory_and_cannot_download(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            trace = _trace(root, include_llm=True)
            verification = root / "verification.json"
            verification.write_text(json.dumps(trace, ensure_ascii=False), encoding="utf-8")
            # The evidence sidecar is present, but no /decisions submission happened.
            write_evidence_for_verification(
                verification,
                {
                    "mode": "candidate_explain",
                    "status": "success",
                    "capabilities": ["candidate", "explain"],
                    "policy_version": "policy-local",
                    "schema_version": "schema-local",
                    "rule_version": "normalized-v2",
                    "trace_id": "trace-local",
                    "job_id": "job-local",
                    "candidates": [],
                    "explanations": [],
                    "discipline_suggestions": [],
                    "calls": [],
                    "warnings": [],
                },
            )
            feedback = root / "feedback.jsonl"
            self.assertFalse(feedback.exists())
            self.assertEqual(memory_hints([_diff()], feedback), {})

            job = Job(
                status="review",
                out_dir=str(root),
                trace_path=str(verification),
                llm_status="success",
                llm_suggestion_count=1,
            )
            store.jobs[job.job_id] = job
            try:
                with self.assertRaises(HTTPException) as raised:
                    asyncio.run(download_result(job.job_id, "draft"))
                self.assertEqual(raised.exception.status_code, 409)
            finally:
                store.jobs.pop(job.job_id, None)

    def test_human_verdict_is_the_only_feedback_source_and_ai_fields_are_not_copied(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            trace = _trace(root, include_llm=True)
            feedback = root / "feedback.jsonl"
            result = finalize_review(
                trace,
                {"diff-local-1": {"verdict": "accept", "note": "人工确认"}},
                reviewer="human-reviewer",
                feedback_path=feedback,
            )

            records = load_feedback(feedback)
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0]["source"], "human")
            self.assertEqual(records[0]["verdict"], "accept")
            self.assertEqual(records[0]["note"], "人工确认")
            raw_feedback = feedback.read_text(encoding="utf-8")
            for forbidden in (_AI_REASON, _AI_CANDIDATE, _CERT_PLACEHOLDER, _DRAFT_PLACEHOLDER):
                self.assertNotIn(forbidden, raw_feedback)
            self.assertNotIn("provider_host", raw_feedback)
            self.assertEqual(result["decisions"]["diff-local-1"]["verdict"], "accept")
            self.assertEqual(result["memory_hints"]["diff-local-1"]["accept"], 1)

    def test_llm_marked_feedback_is_ignored_but_legacy_human_format_remains_compatible(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "feedback.jsonl"
            sig = "safety_list|电源|型号|field_mismatch"
            path.write_text(
                "\n".join(
                    [
                        json.dumps(
                            {
                                "source": "llm",
                                "sig": sig,
                                "key": "diff-ai",
                                "verdict": "accept",
                                "note": _AI_REASON,
                            },
                            ensure_ascii=False,
                        ),
                        json.dumps(
                            {
                                "ts": "2026-01-01T00:00:00+00:00",
                                "reviewer": "legacy-human",
                                "document_version": "draft.docx",
                                "key": "diff-human",
                                "sig": sig,
                                "verdict": "ignore",
                                "note": "旧格式人工备注",
                                "diff_type": "field_mismatch",
                                "field": "型号",
                            },
                            ensure_ascii=False,
                        ),
                        "",
                    ]
                ),
                encoding="utf-8",
            )
            records = load_feedback(path)
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0]["verdict"], "ignore")
            hints = memory_hints([_diff()], path)
            self.assertEqual(hints["diff-local-1"]["ignore"], 1)
            self.assertEqual(hints["diff-local-1"]["accept"], 0)
            self.assertNotIn(_AI_REASON, json.dumps(hints, ensure_ascii=False))

    def test_invalid_decision_shape_unknown_key_and_sensitive_note_do_not_write_feedback(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            feedback = root / "feedback.jsonl"
            diff = _diff()

            with self.assertRaises(ValueError):
                normalize_manual_decisions(
                    {diff.key: {"status": "pending", "reason": _AI_REASON}}
                )
            self.assertFalse(feedback.exists())

            with self.assertRaises(ValueError):
                record_feedback(
                    [diff],
                    {"unknown-key": {"verdict": "accept"}},
                    feedback,
                    reviewer="human-reviewer",
                )
            self.assertFalse(feedback.exists())

            with self.assertRaises(ValueError):
                record_feedback(
                    [diff],
                    {diff.key: {"verdict": "accept", "note": _CERT_PLACEHOLDER}},
                    feedback,
                    reviewer="human-reviewer",
                )
            self.assertFalse(feedback.exists())

    def test_memory_endpoint_recomputes_from_human_feedback_not_review_json(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            trace = _trace(root)
            verification = root / "verification.json"
            verification.write_text(json.dumps(trace, ensure_ascii=False), encoding="utf-8")
            feedback = root / "feedback.jsonl"
            feedback.write_text(
                json.dumps(
                    {
                        "source": "llm",
                        "sig": "safety_list|电源|型号|field_mismatch",
                        "key": "diff-ai",
                        "verdict": "accept",
                        "note": _AI_REASON,
                    },
                    ensure_ascii=False,
                )
                + "\n",
                encoding="utf-8",
            )
            (root / "review.json").write_text(
                json.dumps(
                    {"review": {"memory_hints": {"diff-local-1": {"accept": 99}}}},
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            job = Job(status="review", out_dir=str(root), trace_path=str(verification))
            store.jobs[job.job_id] = job
            try:
                response = TestClient(app).get(f"/api/jobs/{job.job_id}/memory")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json(), {"hints": {}})
            finally:
                store.jobs.pop(job.job_id, None)


class WebDecisionBoundaryTests(unittest.TestCase):
    def test_web_decisions_reject_ai_suggestion_shape_and_unknown_key(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            trace = _trace(root)
            verification = root / "verification.json"
            verification.write_text(json.dumps(trace, ensure_ascii=False), encoding="utf-8")
            job = Job(status="review", out_dir=str(root), trace_path=str(verification))
            store.jobs[job.job_id] = job
            client = TestClient(app)
            try:
                ai_shape = client.post(
                    f"/api/jobs/{job.job_id}/decisions",
                    json={"decisions": {"diff-local-1": {"status": "pending", "reason": _AI_REASON}}},
                )
                self.assertEqual(ai_shape.status_code, 400)
                self.assertEqual(job.decisions, {})

                unknown = client.post(
                    f"/api/jobs/{job.job_id}/decisions",
                    json={"decisions": {"unknown-key": {"verdict": "accept"}}},
                )
                self.assertEqual(unknown.status_code, 400)
                self.assertEqual(job.decisions, {})
            finally:
                store.jobs.pop(job.job_id, None)

    def test_candidate_decisions_cannot_be_submitted_as_final_verdicts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            trace = _trace(root)
            verification = root / "verification.json"
            verification.write_text(json.dumps(trace, ensure_ascii=False), encoding="utf-8")
            job = Job(status="review", out_dir=str(root), trace_path=str(verification))
            store.jobs[job.job_id] = job
            try:
                response = TestClient(app).post(
                    f"/api/jobs/{job.job_id}/decisions",
                    json={
                        "decisions": {"diff-local-1": {"verdict": "accept"}},
                        "candidate_decisions": {_AI_CANDIDATE: "accept"},
                    },
                )
                self.assertEqual(response.status_code, 400)
                self.assertEqual(job.decisions, {})
                self.assertFalse((root / "feedback.jsonl").exists())
            finally:
                store.jobs.pop(job.job_id, None)


if __name__ == "__main__":
    unittest.main()
