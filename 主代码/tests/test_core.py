from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from report_verify.comment_writer import write_comments
from report_verify.document_parser import classify_headers, parse_target_tables
from report_verify.field_map import normalize_for_compare
from report_verify.models import Anchor, DocumentRow, LocatedBlock, VerifyDiff
from report_verify.rules import load_rules, same_value
from report_verify.review import finalize_review
from report_verify.verify_differ import compare_blocks
from webapi.app import app
from webapi.jobs import Job, store


class CoreVerificationTests(unittest.TestCase):
    def _row(self, block: str, name: str, model: str, **values: str) -> DocumentRow:
        fields = {"关键件名称": name, "型号": model, **values}
        return DocumentRow(block, 2, fields, {}, list(fields.values()), 1)

    def test_header_fingerprints(self):
        self.assertEqual(
            classify_headers(["序号", "关键件名称", "型号", "认证标准"]),
            "safety_list",
        )
        self.assertEqual(
            classify_headers(["序号", "关键件名称", "型号", "认证标准", "使用/备用"]),
            "emc_list",
        )
        self.assertEqual(
            classify_headers(["序号", "元件", "名称", "型号", "生产厂(全称)", "本次使用/备用"]),
            "emc_list",
        )
        self.assertEqual(
            classify_headers(["序号和名称", "变更前（原）", "变更后（新）"]),
            "change_confirmation",
        )
        self.assertIsNone(classify_headers(["图片", "说明"]))

    def test_default_normalization_ignores_format_only_differences(self):
        default_rules = load_rules()
        self.assertEqual(default_rules["version"], "normalized-v2")
        self.assertTrue(same_value("A  B", "A B", default_rules))
        self.assertTrue(same_value(" ABC ", "abc", default_rules))
        self.assertTrue(same_value("２２０Ｖ（AC）", "220V(AC)", default_rules))
        self.assertTrue(same_value("250V～", "250V~", default_rules))
        self.assertEqual(normalize_for_compare(" ABC ", default_rules), "abc")

    def test_placeholder_values_equal_blank(self):
        default_rules = load_rules()
        self.assertTrue(same_value("/", "", default_rules))
        self.assertTrue(same_value("—", "", default_rules))
        self.assertTrue(same_value("N/A", "", default_rules))
        self.assertTrue(same_value(" / ", "", default_rules))
        self.assertFalse(same_value("/", "P-1", default_rules))

    def test_list_separator_treats_newline_and_semicolon_as_same(self):
        default_rules = load_rules()
        self.assertTrue(same_value(
            "GB/T18380.12-2022\nGB/T18380.13-2022\nGB/T18380.22-2008",
            "GB/T18380.12-2022;GB/T18380.13-2022;GB/T18380.22-2008",
            default_rules))
        self.assertTrue(same_value("A、B、C", "A;B;C", default_rules))
        self.assertTrue(same_value("A; B; C", "A;B;C", default_rules))
        # 顺序不同仍算差异（选项一）
        self.assertFalse(same_value("A;B;C", "C;B;A", default_rules))
        # 列表里真实值不同仍算差异
        self.assertFalse(same_value(
            "GB/T18380.12-2022;GB/T18380.13-2022",
            "GB/T18380.12-2022;GB/T18380.99-2022",
            default_rules))

    def test_internal_space_in_standard_and_model_is_ignored(self):
        default_rules = load_rules()
        self.assertTrue(same_value("GB/T 6346.14-2023", "GB/T6346.14-2023", default_rules))
        self.assertTrue(same_value("S1000 H", "S1000H", default_rules))
        # 真实字符差异仍保留
        self.assertFalse(same_value("GB/T 6346.14-2023", "GB/T 6346.14-2024", default_rules))

    def test_placeholder_difference_not_flagged_in_block_comparison(self):
        cert = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "电源", "P-1", 位号="/"),
        ])
        draft = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "电源", "P-1", 位号=""),
        ])
        diffs = compare_blocks([cert], [draft])  # 默认规则含 placeholder_blank
        self.assertEqual([d for d in diffs if d.field == "位号"], [])

    def test_default_normalization_is_used_by_block_comparison(self):
        cert = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", " 电源 ", "P-1", 认证标准=" GB-1 "),
        ])
        draft = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "电源", "p-1", 认证标准="gb-1"),
        ])
        self.assertEqual(compare_blocks([cert], [draft]), [])

    def test_strict_mode_and_substantive_differences_remain_visible(self):
        strict = {"version": "strict-test", "enabled": []}
        self.assertFalse(same_value("A  B", "A B", strict))
        self.assertFalse(same_value("ABC", "abc", strict))
        self.assertFalse(same_value("250V～", "250V~", strict))
        default_rules = load_rules()
        self.assertFalse(same_value("P-1", "P-2", default_rules))
        self.assertFalse(same_value("220V", "230V", default_rules))
        self.assertFalse(same_value("CQC-1001", "CQC-1002", default_rules))

    def test_field_missing_and_extra_differences(self):
        cert = LocatedBlock("safety_list", 7, rows=[
            self._row("safety_list", "电源", "P-1", 认证标准="GB-1", 备注="C-1"),
            self._row("safety_list", "仅认证资料", "C-ONLY"),
        ])
        draft = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "电源", "P-1", 认证标准="GB-2", 备注="C-1"),
            self._row("safety_list", "仅报告草稿", "D-ONLY"),
        ])
        diffs = compare_blocks([cert], [draft], {"version": "strict-v1", "enabled": []})
        types = {diff.diff_type for diff in diffs}
        self.assertIn("field_mismatch", types)
        self.assertIn("missing_in_draft", types)
        self.assertIn("extra_in_draft", types)
    def test_row_diffs_aggregate_same_and_changed_rows(self):
        cert_rows = [
            self._row("safety_list", "电源", "P-1", 认证标准="GB-1"),
            self._row("safety_list", "电容", "C-1", 认证标准="GB-2"),
        ]
        draft_rows = [
            self._row("safety_list", "电源", "P-1", 认证标准="GB-9"),
            self._row("safety_list", "电容", "C-1", 认证标准="GB-2"),
        ]
        for index, row in enumerate(cert_rows + draft_rows, start=2):
            row.row_idx = index
        row_diffs = []
        diffs = compare_blocks(
            [LocatedBlock("safety_list", 1, rows=cert_rows)],
            [LocatedBlock("safety_list", 1, rows=draft_rows)],
            {"version": "strict-v1", "enabled": []},
            row_diffs=row_diffs,
        )
        by_name = {row.item_name: row for row in row_diffs}
        self.assertEqual(by_name["电源"].row_status, "CHANGED")
        self.assertEqual(by_name["电源"].changed_fields, ["认证标准"])
        self.assertEqual(by_name["电容"].row_status, "SAME")
        self.assertEqual(by_name["电容"].diff_keys, [])
        self.assertEqual(len(by_name["电源"].diff_keys), 1)
        self.assertEqual(by_name["电源"].diff_keys[0], diffs[0].key)
        self.assertEqual(diffs[0].row_key, by_name["电源"].row_key)

    def test_row_diffs_keep_only_sides_and_empty_placeholder(self):
        cert = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "仅认证", "C-ONLY"),
            self._row("safety_list", "共同", "C-1"),
        ])
        draft = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "共同", "C-1"),
            self._row("safety_list", "仅草稿", "D-ONLY"),
        ])
        for index, row in enumerate(cert.rows + draft.rows, start=2):
            row.row_idx = index
        row_diffs = []
        compare_blocks(
            [cert], [draft], {"version": "strict-v1", "enabled": []},
            row_diffs=row_diffs,
        )
        by_status = {row.row_status: row for row in row_diffs}
        self.assertEqual(by_status["ONLY_CERT"].draft_row, {})
        self.assertEqual(by_status["ONLY_DRAFT"].cert_row, {})
        self.assertTrue(by_status["ONLY_CERT"].diff_keys)
        self.assertTrue(by_status["ONLY_DRAFT"].diff_keys)

    def test_unassigned_draft_table_has_row_diff(self):
        draft = LocatedBlock("safety_list", 9, header=["序号", "关键件名称", "型号"], rows=[
            self._row("safety_list", "DRAFT 独有表行", "D-ONLY"),
        ])
        row_diffs = []
        diffs = compare_blocks(
            [], [draft], {"version": "strict-v1", "enabled": []},
            row_diffs=row_diffs,
        )
        table_diff = next(diff for diff in diffs if diff.diff_type == "table_extra_in_draft")
        self.assertEqual(len(row_diffs), 1)
        self.assertEqual(row_diffs[0].row_status, "ONLY_DRAFT")
        self.assertEqual(row_diffs[0].row_kind, "block")
        self.assertEqual(row_diffs[0].diff_keys, [table_diff.key])

    def test_row_diff_api_exposes_decision_state(self):
        self.assertEqual(TestClient(app).get("/api/health").json(), {"status": "ok"})
        with tempfile.TemporaryDirectory() as directory:
            trace_path = Path(directory) / "verification.json"
            trace_path.write_text(json.dumps({
                "result": {"row_diffs": [{
                    "row_key": "row-test", "block_type": "safety_list",
                    "row_status": "CHANGED", "row_kind": "row", "item_name": "电源",
                    "cert_row": {"型号": "P-2"}, "draft_row": {"型号": "P-1"},
                    "cert_row_ref": {}, "draft_row_ref": {},
                    "changed_fields": ["型号"], "diff_keys": ["diff-test"],
                    "match_confidence": 1.0, "match_reason": "测试",
                    "match_status": "matched",
                }]}
            }, ensure_ascii=False), encoding="utf-8")
            job = Job(status="review", trace_path=str(trace_path))
            job.decisions = {"diff-test": {"verdict": "accept", "note": "已确认"}}
            store.jobs[job.job_id] = job
            try:
                response = TestClient(app).get(f"/api/jobs/{job.job_id}/row-diffs")
                self.assertEqual(response.status_code, 200)
                payload = response.json()
                self.assertEqual(payload["total"], 1)
                self.assertEqual(payload["items"][0]["row_key"], "row-test")
                self.assertEqual(payload["items"][0]["review_status"], "accept")
                self.assertEqual(payload["items"][0]["decisions"]["diff-test"]["note"], "已确认")
            finally:
                store.jobs.pop(job.job_id, None)

    def test_vertical_merge_carries_name_and_serial(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "merged.docx"
            document = Document()
            table = document.add_table(rows=3, cols=4)
            for cell, value in zip(table.rows[0].cells, ["序号", "关键件名称", "型号", "认证标准"]):
                cell.text = value
            table.cell(1, 0).merge(table.cell(2, 0)).text = "1"
            table.cell(1, 1).merge(table.cell(2, 1)).text = "电源"
            table.cell(1, 2).text = "P-1"
            table.cell(1, 3).text = "GB-1"
            table.cell(2, 2).text = "P-2"
            table.cell(2, 3).text = "GB-1"
            document.save(source)
            blocks, _ = parse_target_tables(str(source), role="DRAFT")
            self.assertEqual(len(blocks), 1)
            self.assertEqual([row.values["关键件名称"] for row in blocks[0].rows], ["电源", "电源"])
            self.assertEqual([row.values["序号"] for row in blocks[0].rows], ["1", "1"])

    def test_consolidate_duplicate_certification_column(self):
        from report_verify.document_parser import _consolidate_duplicate_fields
        from report_verify.models import Anchor
        values = {
            "位号": "CX802,CX803,CX804", "位号__2": "/",
            "关键件名称": "X电容", "型号": "CMPP",
            "制造商": "威海汉城成镐电子有限公司", "生产企业": "威海汉城成镐电子有限公司",
            "认证标准": "威海汉城成镐电子有限公司", "认证标准__2": "GB/T6346.14-2023",
            "备注": "CQC13001098793", "证书号": "CQC13001098793",
        }
        anchors = {k: Anchor(7, 199, i, v[:80]) for i, (k, v) in enumerate(values.items())}
        _consolidate_duplicate_fields(values, anchors)
        self.assertEqual(values["认证标准"], "GB/T6346.14-2023")
        self.assertEqual(values["位号"], "CX802,CX803,CX804")
        # __2 列原样保留供追溯
        self.assertEqual(values["认证标准__2"], "GB/T6346.14-2023")
        # 认证标准锚点跟随被选中的 __2 列，指向真实标准号所在网格
        self.assertEqual(anchors["认证标准"], anchors["认证标准__2"])

    def test_consolidate_leaves_normal_rows_unchanged(self):
        from report_verify.document_parser import _consolidate_duplicate_fields
        values = {
            "位号": "CX1", "关键件名称": "电源", "型号": "P-1",
            "认证标准": "GB-1", "证书号": "CERT-1",
        }
        anchors: dict = {}
        original = dict(values)
        _consolidate_duplicate_fields(values, anchors)
        self.assertEqual(values, original)

    def test_parse_duplicate_certification_column_from_docx(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "dup_cert.docx"
            document = Document()
            headers = ["序号", "关键件名称", "型号", "生产企业", "认证标准", "认证标准"]
            table = document.add_table(rows=2, cols=len(headers))
            for cell, value in zip(table.rows[0].cells, headers):
                cell.text = value
            data = table.rows[1]
            data.cells[0].text = "1"
            data.cells[1].text = "X电容"
            data.cells[2].text = "CMPP"
            # 生产企业与第一个认证标准列横向合并，公司名溢出到 idx4
            data.cells[3].merge(data.cells[4]).text = "威海汉城成镐电子有限公司"
            data.cells[5].text = "GB/T6346.14-2023"
            document.save(source)
            blocks, _ = parse_target_tables(str(source), role="CERT")
            self.assertEqual(len(blocks), 1)
            row = blocks[0].rows[0]
            self.assertEqual(row.values["认证标准"], "GB/T6346.14-2023")
            self.assertNotEqual(row.values["认证标准"], "威海汉城成镐电子有限公司")

    def test_model_subset_is_matched_and_model_difference_is_visible(self):
        cert = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "电源", "P-1、P-2", 认证标准="GB-1"),
        ])
        draft = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "电源", "P-1", 认证标准="GB-1"),
        ])
        diffs = compare_blocks([cert], [draft], {"version": "strict-v1", "enabled": []})
        self.assertEqual({diff.diff_type for diff in diffs}, {"field_mismatch"})
        self.assertEqual(next(diff for diff in diffs if diff.field == "型号").match_status, "matched")

    def test_model_overlap_is_matched_as_needs_review(self):
        cert = LocatedBlock("safety_list", 1, rows=[self._row("safety_list", "电源", "P-1、P-2")])
        draft = LocatedBlock("safety_list", 1, rows=[self._row("safety_list", "电源", "P-2、P-3")])
        diffs = compare_blocks([cert], [draft], {"version": "strict-v1", "enabled": []})
        model_diff = next(diff for diff in diffs if diff.field == "型号")
        self.assertEqual(model_diff.match_status, "needs_review")
        self.assertEqual({diff.diff_type for diff in diffs}, {"field_mismatch"})

    def test_empty_model_does_not_consume_a_modelled_row(self):
        cert = LocatedBlock("safety_list", 1, rows=[self._row("safety_list", "电源", "")])
        draft = LocatedBlock("safety_list", 1, rows=[self._row("safety_list", "电源", "P-1")])
        diffs = compare_blocks([cert], [draft], {"version": "strict-v1", "enabled": []})
        self.assertEqual({diff.diff_type for diff in diffs}, {"missing_in_draft", "extra_in_draft"})

    def test_duplicate_rows_are_consumed_one_to_one(self):
        cert = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "电容", "C-1", 规格材料="10uF"),
            self._row("safety_list", "电容", "C-1", 规格材料="20uF"),
        ])
        draft = LocatedBlock("safety_list", 1, rows=[
            self._row("safety_list", "电容", "C-1", 规格材料="10uF"),
            self._row("safety_list", "电容", "C-1", 规格材料="30uF"),
        ])
        diffs = compare_blocks([cert], [draft], {"version": "strict-v1", "enabled": []})
        self.assertEqual(len([diff for diff in diffs if diff.diff_type == "field_mismatch"]), 1)
        self.assertNotIn("missing_in_draft", {diff.diff_type for diff in diffs})
        self.assertNotIn("extra_in_draft", {diff.diff_type for diff in diffs})

    def test_nonmatching_groups_cannot_cross_match(self):
        cert_row = self._row("safety_list", "电源", "P-1")
        draft_row = self._row("safety_list", "电源", "P-1")
        cert_row.group = "13062-D型机内电源单元"
        draft_row.group = "55431-F型机内电源单元"
        diffs = compare_blocks(
            [LocatedBlock("safety_list", 1, rows=[cert_row])],
            [LocatedBlock("safety_list", 1, rows=[draft_row])],
            {"version": "strict-v1", "enabled": []},
        )
        self.assertEqual({diff.diff_type for diff in diffs}, {"missing_in_draft", "extra_in_draft"})

    def test_different_names_with_same_model_and_certificate_are_paired(self):
        cert = LocatedBlock("safety_list", 1, header=["序号", "关键件名称", "型号", "证书号"], rows=[
            self._row("safety_list", "隔离电容器（内置）", "C-1", 证书号="CERT-1"),
        ])
        draft = LocatedBlock("safety_list", 1, header=["序号", "关键件名称", "型号", "证书号"], rows=[
            self._row("safety_list", "天线隔离电容(内置)", "C-1", 证书号="CERT-1"),
        ])
        diffs = compare_blocks([cert], [draft], {"version": "strict-v1", "enabled": []})
        self.assertEqual({diff.diff_type for diff in diffs}, {"field_mismatch"})
        name_diff = next(diff for diff in diffs if diff.field == "关键件名称")
        self.assertEqual(name_diff.match_status, "matched")
        self.assertIn("名称不同但型号与证书/规格身份强一致", name_diff.match_reason)

    def test_same_table_short_name_and_wildcard_model_are_paired_but_visible(self):
        cert = LocatedBlock("emc_list", 28, header=["序号", "关键件名称", "型号"], rows=[
            self._row("emc_list", "主板", "RSAG7.820.13199******"),
        ])
        draft = LocatedBlock("emc_list", 2, header=["序号", "关键件名称", "型号"], rows=[
            self._row("emc_list", "主板(含网卡、含调谐单元)", "RSAG7.820.13199*"),
        ])
        diffs = compare_blocks([cert], [draft], {"version": "strict-v1", "enabled": []})
        self.assertEqual({diff.diff_type for diff in diffs}, {"field_mismatch"})
        self.assertEqual({diff.field for diff in diffs}, {"关键件名称", "型号"})
        self.assertTrue(all(diff.match_status == "needs_review" for diff in diffs))
        model_diff = next(diff for diff in diffs if diff.field == "型号")
        self.assertIn("通配符族相同", model_diff.match_reason)

    def test_missing_optional_block_is_not_expanded_to_row_missing(self):
        cert = LocatedBlock("product_description", 5, rows=[
            DocumentRow("product_description", 2, {"项目": "额定值", "值": "220V"}),
        ])
        diffs = compare_blocks([cert], [], {"version": "strict-v1", "enabled": []})
        self.assertEqual(len(diffs), 1)
        self.assertEqual(diffs[0].diff_type, "block_missing_in_draft")
        self.assertEqual(diffs[0].match_status, "needs_review")

    def test_comment_writer_creates_comments_xml(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "draft.docx"
            target = Path(directory) / "annotated.docx"
            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "关键件名称"
            table.cell(0, 1).text = "型号"
            table.cell(1, 0).text = "电源"
            table.cell(1, 1).text = "P-1"
            document.save(source)
            diff = VerifyDiff(
                key="test-1", block_type="safety_list", item_name="电源",
                diff_type="field_mismatch", field="型号", cert_value="P-2",
                draft_value="P-1", draft_anchor=Anchor(1, 2, 1, "P-1"),
            )
            report = write_comments(str(source), [diff], str(target))
            self.assertEqual(report["written"], 1)
            self.assertTrue(diff.comment_written)
            with ZipFile(target) as package:
                self.assertIn("word/comments.xml", package.namelist())
    def test_review_generates_final_outputs_only_after_all_decisions(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            root = Path(directory)
            draft = root / "draft.docx"
            source = Document()
            table = source.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "字段"
            table.cell(1, 0).text = "旧值"
            source.save(draft)
            diff = VerifyDiff(
                key="review-1", block_type="safety_list", item_name="项目",
                diff_type="field_mismatch", field="型号", cert_value="新值",
                draft_value="旧值", draft_anchor=Anchor(1, 2, 0, "旧值"),
            )
            trace = {
                "cert_path": "A20-cert.docx", "draft_path": "100-draft.docx",
                "draft_docx": str(draft),
                "outputs": {"annotated_draft": str(root / "initial.docx")},
                "result": {"diffs": [json.loads(json.dumps({
                    "key": diff.key, "block_type": diff.block_type,
                    "item_name": diff.item_name, "diff_type": diff.diff_type,
                    "field": diff.field, "cert_value": diff.cert_value,
                    "draft_value": diff.draft_value, "cert_location": "表1第2行",
                    "draft_anchor": {"table_idx": 1, "row_idx": 2, "cell_idx": 0, "para_hint": "旧值"},
                    "discipline": None, "discipline_status": "discipline_pending",
                    "rule_version": "strict-v1", "comment_written": False,
                    "comment_error": "", "review_status": "unreviewed",
                    "review_note": "", "reviewed_by": "",
                }))], "warnings": []},
            }
            result = finalize_review(
                trace, {"review-1": {"verdict": "accept", "note": "确认修改"}},
                reviewer="tester", feedback_path=root / "feedback.jsonl",
            )
            self.assertEqual(result["status"], "completed")
            self.assertTrue(Path(result["outputs"]["annotated_draft"]).exists())
            self.assertTrue(Path(result["outputs"]["summary"]).exists())
            self.assertTrue((root / "feedback.jsonl").exists())


    def test_100_is_draft_and_a20_is_cert(self):
        sample_dir = Path(__file__).resolve().parents[2] / "确认报告草稿样本"
        draft_candidates = [p for p in (sample_dir / "_converted").glob("*.docx") if p.name.startswith("100")]
        cert_candidates = [p for p in sample_dir.glob("*.docx") if p.name.startswith("A20")]
        self.assertTrue(draft_candidates, "没有找到 100 开头的 DRAFT 样本")
        self.assertTrue(cert_candidates, "没有找到 A20 开头的 CERT 样本")
        self.assertTrue(draft_candidates[0].name.startswith("100"))
        self.assertTrue(cert_candidates[0].name.startswith("A20"))


if __name__ == "__main__":
    unittest.main()
