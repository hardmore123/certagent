"""Route 2: convert legacy .doc to .docx, then parse with python-docx.

The Word conversion is deliberately reported separately from the python-docx
parse. This lets the benchmark distinguish conversion risk from parser quality.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx.table import _Cell

from .common import clean, field_for_header, make_row, row_dicts

_CARRY = ["序号", "关键件名称"]


def convert_to_docx(source: Path, output_dir: Path) -> Path:
    """Convert a legacy binary Word document using local Word COM only."""
    if source.suffix.lower() == ".docx":
        return source
    try:
        import win32com.client as win32  # type: ignore
    except Exception as exc:
        raise RuntimeError("docx_route 需要本机 Microsoft Word/pywin32 完成 .doc 转换") from exc

    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / f"{source.stem}.converted.docx"
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    doc = None
    try:
        doc = word.Documents.Open(
            str(source.resolve()), ReadOnly=True, AddToRecentFiles=False
        )
        try:
            doc.SaveAs2(str(target), FileFormat=16)
        except Exception:
            doc.SaveAs(str(target), FileFormat=16)
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass
        try:
            word.Quit()
        except Exception:
            pass
    if not target.exists():
        raise RuntimeError(f"Word 转换未生成目标文件: {target}")
    return target


def _span(tc: Any) -> int:
    try:
        value = tc.tcPr.gridSpan.val
        return max(int(value), 1) if value else 1
    except Exception:
        return 1


def _is_vertical_merge(tc: Any) -> bool:
    try:
        return tc.tcPr.vMerge is not None
    except Exception:
        return False


def _xml_cells(row: Any, table: Any) -> list[tuple[int, str, int, bool]]:
    """Return logical start column, text, grid span and vertical-merge flag."""
    cells: list[tuple[int, str, int, bool]] = []
    column = 1
    for tc in row._tr.tc_lst:
        span = _span(tc)
        # CT_Tc.text is None in recent python-docx; _Cell exposes paragraph text.
        cells.append((column, clean(_Cell(tc, table).text), span, _is_vertical_merge(tc)))
        column += span
    return cells


def _header_map(cells: list[tuple[int, str, int, bool]]) -> dict[int, str]:
    return {
        column: field
        for column, text, _span_value, _vertical in cells
        if (field := field_for_header(text))
    }


def _parse_table(table: Any, table_no: int) -> tuple[list[Any], list[str]]:
    rows: list[Any] = []
    group_events: list[str] = []
    if not table.rows:
        return rows, group_events

    header_cells = _xml_cells(table.rows[0], table)
    col_map = _header_map(header_cells)
    header_max = max((column for column, _text, _span_value, _vertical in header_cells), default=0)
    model_ci = next((ci for ci, field in col_map.items() if field == "型号"), None)
    merge_point = model_ci + 1 if model_ci is not None else None
    carry: dict[str, str] = {}
    merged_values: dict[str, str] = {}

    def field_for(column: int, row_max: int) -> str | None:
        if merge_point and header_max and row_max < header_max and column >= merge_point:
            direct = col_map.get(column + 1)
            if direct:
                return direct
        direct = col_map.get(column)
        if direct:
            return direct
        # Word 转换后，某些合并行会把单元格起始列落在两个表头列之间。
        # 将该单元格归到右侧最近的已知表头，避免制造商/生产企业错位。
        for header_column in sorted(col_map):
            if header_column > column:
                return col_map[header_column]
        return None

    for row_index, word_row in enumerate(table.rows[1:], start=2):
        cells = _xml_cells(word_row, table)
        non_empty = [text for _column, text, _span_value, _vertical in cells if text]
        if len(cells) == 1 and non_empty:
            carry["group"] = non_empty[0]
            group_events.append(non_empty[0])
            continue
        if not non_empty:
            continue

        row_max = max((column for column, _text, _span_value, _vertical in cells), default=0)
        data: dict[str, Any] = {
            "table": table_no,
            "group": carry.get("group", ""),
            "row_index": row_index,
        }
        raw: dict[str, str] = {}
        for column, text, _span_value, vertical in cells:
            raw[str(column)] = text
            field = field_for(column, row_max)
            if not field:
                continue
            if text:
                data[field] = text
                if vertical:
                    merged_values[field] = text
            elif vertical and field in merged_values:
                # python-docx leaves continuation cells empty for w:vMerge;
                # Word COM exposes the merged value on each logical row.
                data[field] = merged_values[field]
        if table_no == 1 and data.get("备注"):
            data["证书号"] = data["备注"]
        for field in _CARRY:
            if not data.get(field):
                data[field] = carry.get(field, "")
            else:
                carry[field] = data[field]
        rows.append(make_row(data, raw))
    return rows, group_events


def parse_docx(path: Path) -> tuple[list[Any], dict[str, Any]]:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("未安装 python-docx，无法执行 docx_route") from exc

    document = Document(str(path))
    rows: list[Any] = []
    groups: list[str] = []
    for table_no, table in enumerate(document.tables, start=1):
        table_rows, table_groups = _parse_table(table, table_no)
        rows.extend(table_rows)
        groups.extend(table_groups)
    return rows, {
        "table_count": len(document.tables),
        "group_events": groups,
        "group_event_count": len(groups),
    }


def run_docx_route(source_path: str, work_dir: str) -> dict[str, Any]:
    """Run conversion and parsing, returning benchmark-friendly JSON data."""
    source = Path(source_path).resolve()
    work = Path(work_dir).resolve()
    converted = convert_to_docx(source, work / "converted")
    rows, metadata = parse_docx(converted)
    return {
        "status": "success",
        "rows": row_dicts(rows),
        "metadata": {
            "source": str(source),
            "converted_docx": str(converted),
            **metadata,
        },
    }
